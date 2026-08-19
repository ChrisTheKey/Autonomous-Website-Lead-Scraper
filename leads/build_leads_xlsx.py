# -*- coding: utf-8 -*-
import sys, os
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lead_data import L
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

COLS = ["Lead ID","Priorität","Lead Score","Firmenname","Vorname Entscheider","Nachname Entscheider",
"Position Entscheider","Ort","Adresse","Telefonnummer","Direkte Telefonnummer Entscheider","E-Mail",
"Website","Quelle Entscheider","Quelle Kontaktdaten","Dienstleistungsangebot","Website Score",
"Konkrete Website-Schwächen","Grösste Verkaufschance","Empfohlenes Website-Paket","Website-Listenpreis CHF",
"Potenzielle Agents","Begründung Agents","Agent-Upsell-Potenzial","Google-Bewertungen / Reputation Hinweis",
"Besonderer Gesprächsaufhänger","Call-Status","Entscheider erreicht?","Datum letzter Anruf","Hat abgenommen?",
"Ergebnis Opener-Call","Haupteinwand","Closer-Call vereinbart?","Closer-Call Datum","Closer-Call Uhrzeit",
"Zoom-Link gesendet?","Closing Probability %","Nächster Schritt","Follow-up Datum","Anzahl Kontaktversuche",
"Verkaufsstatus","Notizen"]

for d in L:
    d["lead_score"] = d["b_web"] + d["b_pot"] + d["b_err"] + d["b_ag"]
    s = d["lead_score"]
    d["prio"] = "A" if s >= 8 else ("B" if s >= 6 else "C")

# genau 50: die zwei schwächsten Leads entfernen
L.sort(key=lambda d: (-d["lead_score"], d["ws"], d["name"]))
dropped = L[50:]
L[:] = L[:50]

PRIO_ORDER = {"A":0,"B":1,"C":2}
L.sort(key=lambda d: (PRIO_ORDER[d["prio"]], -d["lead_score"], d["ws"], d["name"]))

wb = Workbook(); ws = wb.active; ws.title = "Leads Bern"

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PRIO_FILL = {"A": PatternFill("solid", fgColor="C6EFCE"),
             "B": PatternFill("solid", fgColor="FFEB9C"),
             "C": PatternFill("solid", fgColor="FFC7CE")}
SALES_FILL = PatternFill("solid", fgColor="F2F2F2")

ws.append(COLS)
for c in ws[1]:
    c.fill = HDR_FILL; c.font = HDR_FONT
    c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
ws.freeze_panes = "D2"

SALES_COLS = set(range(27, 43))  # 1-basiert: Call-Status .. Notizen

for i, d in enumerate(L, start=1):
    row = [
        f"BE-{i:03d}", d["prio"], d["lead_score"], d["name"], d["vn"], d["nn"], d["pos"],
        d["ort"], d["adresse"], d["tel"], d["teldirekt"], d["mail"], d["web"],
        d["qentscheider"], d["qkontakt"], d["services"], d["ws"], d["schwaechen"], d["chance"],
        d["paket"], d["preis"], d["agents"], d["agentbegr"],
        ("Ja" if d["b_ag"] == 1 else "Nein"), d["reputation"], d["aufhaenger"],
        "Noch nicht angerufen", "Offen", "", "Offen", "", "", "Nein", "", "", "", "",
        "Erstkontakt Opener-Call", "", 0, "Neuer Lead", "",
    ]
    ws.append(row)

for r in range(2, ws.max_row + 1):
    prio = ws.cell(r, 2).value
    for c in range(1, len(COLS) + 1):
        cell = ws.cell(r, c)
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=(c in (4,7,8,9,14,15,16,18,19,22,23,25,26,42)))
        if c in SALES_COLS: cell.fill = SALES_FILL
    ws.cell(r, 2).fill = PRIO_FILL[prio]
    ws.cell(r, 2).font = Font(bold=True)
    ws.cell(r, 2).alignment = Alignment(horizontal="center", vertical="top")
    ws.cell(r, 1).font = Font(bold=True)
    for c in (3, 17, 21, 40):
        ws.cell(r, c).alignment = Alignment(horizontal="center", vertical="top")
    ws.cell(r, 21).number_format = '#,##0'
    ws.cell(r, 13).font = Font(color="0563C1", underline="single")
    ws.cell(r, 29).number_format = 'DD.MM.YYYY'
    ws.cell(r, 34).number_format = 'DD.MM.YYYY'
    ws.cell(r, 39).number_format = 'DD.MM.YYYY'
    ws.cell(r, 37).number_format = '0"%"'

WIDTHS = {1:10,2:9,3:11,4:38,5:18,6:18,7:30,8:22,9:44,10:16,11:34,12:30,13:40,14:46,15:34,16:52,
          17:13,18:78,19:56,20:22,21:15,22:34,23:78,24:12,25:52,26:60,27:20,28:18,29:18,30:16,
          31:26,32:22,33:20,34:18,35:20,36:20,37:20,38:26,39:16,40:12,41:18,42:34}
for c, w in WIDTHS.items(): ws.column_dimensions[get_column_letter(c)].width = w
ws.row_dimensions[1].height = 42
for r in range(2, ws.max_row + 1): ws.row_dimensions[r].height = 120

t = Table(displayName="LeadsBern", ref=f"A1:{get_column_letter(len(COLS))}{ws.max_row}")
t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
ws.add_table(t)
ws.auto_filter.ref = t.ref

out = "/home/user/Autonomous-Website-Lead-Scraper/leads/ZERO_Umzugsfirmen_Leads_Bern.xlsx"
os.makedirs(os.path.dirname(out), exist_ok=True)
wb.save(out)

a = sum(1 for d in L if d["prio"] == "A"); b = sum(1 for d in L if d["prio"] == "B"); c = sum(1 for d in L if d["prio"] == "C")
avg = sum(d["ws"] for d in L) / len(L)
print(f"Leads: {len(L)} | A={a} B={b} C={c} | Ø Website Score = {avg:.2f}")
print("Entfernt (Überhang):", ", ".join(x["name"] for x in dropped) or "-")
print("Ohne Entscheidername:", sum(1 for d in L if d["vn"]=="Nicht gefunden" and d["nn"]=="Nicht gefunden"))
print("Ohne E-Mail:", sum(1 for d in L if d["mail"].startswith("Nicht gefunden")))
print("Datei:", out)
